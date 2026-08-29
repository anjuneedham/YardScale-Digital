"""Rasterise a .pptx with python-pptx + PIL for visual QA.

LibreOffice cannot load files in this container, so this stands in for
soffice/pdftoppm. Liberation Sans is metric-compatible with Arial and
Liberation Mono with Courier New; Calibri is approximated with Liberation
Sans, which is wider - so any text-fit check here is conservative.
"""
import io
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

DPI = 110
FONTS = {
    ('sans', False, False): '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
    ('sans', True, False): '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
    ('sans', False, True): '/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf',
    ('sans', True, True): '/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf',
    ('mono', False, False): '/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf',
    ('mono', True, False): '/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf',
    ('mono', False, True): '/usr/share/fonts/truetype/liberation/LiberationMono-Italic.ttf',
    ('mono', True, True): '/usr/share/fonts/truetype/liberation/LiberationMono-BoldItalic.ttf',
}
_cache = {}


def font(face, size_pt, bold, italic):
    fam = 'mono' if 'courier' in (face or '').lower() else 'sans'
    px = max(6, int(round(size_pt * DPI / 72.0)))
    key = (fam, bold, italic, px)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTS[(fam, bold, italic)], px)
    return _cache[key]


def px(emu):
    return int(round(Emu(emu).inches * DPI))


def rgb(color, default=None):
    try:
        if color and color.type is not None and color.rgb is not None:
            return tuple(bytes.fromhex(str(color.rgb)))
    except Exception:
        pass
    return default


overflow_report = []


def draw_text_frame(draw, shape, box, slide_no):
    x0, y0, w, h = box
    pad_l = pad_r = pad_t = 0
    try:
        tf = shape.text_frame
        pad_l = px(tf.margin_left or 0)
        pad_r = px(tf.margin_right or 0)
        pad_t = px(tf.margin_top or 0)
    except Exception:
        return
    avail = max(10, w - pad_l - pad_r)

    lines = []          # (text, font, color, align, height)
    for para in tf.paragraphs:
        runs = [r for r in para.runs]
        if not runs:
            lines.append(('', font('sans', 12, False, False), (0, 0, 0), para.alignment, int(12 * DPI / 72)))
            continue
        align = para.alignment
        # wrap the paragraph run-by-run, carrying style per fragment
        cur, cur_w = [], 0
        for r in runs:
            f = r.font
            sz = f.size.pt if f.size else 18
            fo = font(f.name or 'Calibri', sz, bool(f.bold), bool(f.italic))
            col = rgb(f.color, (0, 0, 0))
            for word in r.text.split(' '):
                if word == '':
                    continue
                token = word + ' '
                tw = draw.textlength(token, font=fo)
                if cur and cur_w + tw > avail:
                    lines.append((cur, None, None, align, max(fr[3] for fr in cur)))
                    cur, cur_w = [], 0
                cur.append((token, fo, col, int(fo.size * 1.22)))
                cur_w += tw
            if '\n' in r.text or r is runs[-1]:
                pass
        if cur:
            lines.append((cur, None, None, align, max(fr[3] for fr in cur)))

    total = sum(l[4] for l in lines)
    if total > h + 4 and total > 0:
        overflow_report.append(
            f'  slide {slide_no}: text may overflow — need {total}px, box {h}px :: '
            + repr(shape.text_frame.text[:70])
        )

    # vertical alignment
    va = None
    try:
        va = shape.text_frame.vertical_anchor
    except Exception:
        pass
    y = y0 + pad_t
    if str(va) and 'MIDDLE' in str(va):
        y = y0 + max(0, (h - total) // 2)
    elif str(va) and 'BOTTOM' in str(va):
        y = y0 + max(0, h - total)

    for frags, _f, _c, align, lh in lines:
        if not frags:
            y += lh
            continue
        line_w = sum(draw.textlength(t, font=f) for t, f, _, _ in frags)
        if str(align) and 'CENTER' in str(align):
            x = x0 + pad_l + max(0, (avail - line_w) / 2)
        elif str(align) and 'RIGHT' in str(align):
            x = x0 + pad_l + max(0, avail - line_w)
        else:
            x = x0 + pad_l
        for t, f, c, _ in frags:
            draw.text((x, y), t, font=f, fill=c)
            x += draw.textlength(t, font=f)
        y += lh


def render(path, prefix):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        bg = (255, 255, 255)
        try:
            bg = rgb(slide.background.fill.fore_color, (255, 255, 255))
        except Exception:
            pass
        im = Image.new('RGB', (W, H), bg)
        d = ImageDraw.Draw(im)
        for sh in slide.shapes:
            if sh.left is None:
                continue
            box = (px(sh.left), px(sh.top), px(sh.width), px(sh.height))
            x, y, w, h = box
            st = str(sh.shape_type)
            if 'PICTURE' in st:
                try:
                    pic = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA')
                    pic = pic.resize((max(1, w), max(1, h)))
                    im.paste(pic, (x, y), pic)
                except Exception as e:
                    d.rectangle([x, y, x + w, y + h], outline=(255, 0, 0))
                continue
            fill = None
            try:
                if sh.fill.type is not None and str(sh.fill.type) != 'MSO_FILL_TYPE.BACKGROUND (5)':
                    fill = rgb(sh.fill.fore_color)
            except Exception:
                pass
            line = None
            try:
                line = rgb(sh.line.color)
            except Exception:
                pass
            if fill or line:
                shp = getattr(sh, 'shape_type', None)
                if shp is not None and 'OVAL' in str(shp):
                    d.ellipse([x, y, x + w, y + h], fill=fill, outline=line)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=fill, outline=line)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(d, sh, box, i)
        f = f'{prefix}-{i}.png'
        im.save(f)
        out.append(f)
    return out


if __name__ == '__main__':
    files = render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else 'qa')
    print('\n'.join(files))
    if overflow_report:
        print('\nPOSSIBLE TEXT OVERFLOW (conservative — Liberation is wider than Calibri):')
        print('\n'.join(overflow_report))
    else:
        print('\nno text-fit warnings')
